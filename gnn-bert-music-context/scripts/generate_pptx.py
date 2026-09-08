#!/usr/bin/env python3
"""Build a standalone PowerPoint of the 12-slide course talk.

    python scripts/generate_pptx.py
"""

from __future__ import annotations

import shutil
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]

INK = RGBColor(0x10, 0x0E, 0x0C)
PANEL = RGBColor(0x17, 0x14, 0x10)
CREAM = RGBColor(0xEF, 0xE6, 0xD6)
MUTE = RGBColor(0xB7, 0xAA, 0x98)
COPPER = RGBColor(0xC4, 0x5C, 0x26)
LINE = RGBColor(0x3A, 0x32, 0x28)

SLIDES = [
    {
        "kicker": "Course presentation",
        "title": "GNN-Based BERT for Understanding Context from Music",
        "layout": "title",
        "bullets": [
            "Understanding, not generation",
            "Official FMA-small subset · 511 clips · 386 artists",
            "Four tasks: BERT tags, GraphSAGE, fusion, InfoNCE",
        ],
        "footnote": "Cadence Lab · CSE425 / EEE474 / CSE715 · September 2026",
        "notes": (
            "This is Cadence Lab's course project for Neural Networks: GNN-based BERT for "
            "understanding context from music. The assignment is not to generate audio. It is "
            "to read a clip as structure plus language, then predict tags, genre, mood, and "
            "retrieval. Everything I will show was trained on official Free Music Archive small "
            "audio: five hundred eleven thirty-second clips, artist-disjoint splits. I will be "
            "honest about where the labels leak from metadata, because that is the actual "
            "finding, not a leaderboard claim. After the numbers I will point you to the lab, "
            "the saved graphs, and the seven-page IEEE report."
        ),
    },
    {
        "kicker": "Motivation",
        "title": "A clip is harmonic, timbral, and linguistic at once",
        "layout": "split",
        "split": [
            ("CNN on log-mel", ["Local timbre and transients", "Blind to chord loops", "Blind to repeated segments"]),
            ("LM on captions", ["Instruments, genre, mood words", "No alignment to time", "Cannot hear a drop"]),
            ("This project", ["Graph G from the spectrogram", "TinyBERT on the caption", "Fuse, then retrieve"]),
        ],
        "notes": (
            "A recording can be folk, acoustic, guitar-led, and built from a short chord loop "
            "at the same time. A spectrogram CNN hears local timbre. It does not see that bar "
            "three repeats bar one, or that A minor goes to C major. A language model on a "
            "caption hears a folk recording titled Arabesque by Ed Askew. It does not know "
            "when those words line up with a modulation. The brief therefore asks for a hybrid: "
            "BERT on text, a graph net on a structure graph from audio, a fusion stage, and a "
            "contrastive retriever. Five beats in this talk: why, representation, the four "
            "models, held-out numbers, then what those numbers actually mean."
        ),
    },
    {
        "kicker": "Assignment",
        "title": "Four tasks, one pipeline",
        "layout": "split",
        "split": [
            ("1  BERT tags", ["TinyBERT CLS", "24-way BCE", "Swap-in for bert-base"]),
            ("2  Graphs", ["GraphSAGE / GAT", "Segment + chord G", "CNN control"]),
            ("3  Fusion", ["Concat vs cross-attn", "Tag + VA heads", "Warm-start BERT"]),
            ("4  Retrieval", ["Two towers", "InfoNCE tau = 0.07", "R@1 / 5 / 10"]),
        ],
        "notes": (
            "Task one is easy on paper: a BERT tag classifier on captions, twenty-four "
            "MagnaTagATune-style labels, binary cross-entropy. Task two is the graph: GraphSAGE, "
            "and GAT, on a segment-similarity graph and a chord-transition graph, against a CNN "
            "mel baseline. Task three fuses the graph readout with token states, concat versus "
            "cross-attention, plus valence and arousal heads. Task four is a dual encoder with "
            "InfoNCE. We also ship majority and random tag baselines so the table is complete."
        ),
    },
    {
        "kicker": "Representation",
        "title": "A track is the tuple T = (audio, text, G, y)",
        "layout": "bullets",
        "bullets": [
            "X_audio: log-mel 128, chroma 12, MFCC 13, RMS, centroid, ZCR",
            "X_text: at most 48 tokens, vocabulary rebuilt from FMA captions",
            "G: 30-node segment graph or ~8-node chord graph",
            "y: 24 tags, 8 genres, valence / arousal in [0, 1]",
        ],
        "notes": (
            "Audio is a one-twenty-eight-bin log-mel at twenty-two thousand fifty Hertz, hop five "
            "twelve, plus chroma, MFCCs, RMS, centroid, and zero-crossing rate, all z-scored per "
            "clip. Text is a whitespace-tokenized caption of at most forty-eight tokens with CLS "
            "and SEP. G is either a thirty-node segment graph or a chord-transition graph. Labels "
            "y are twenty-four binary tags, eight-way genre, and valence-arousal in zero to one. "
            "Jazz and metal heads exist in the classifier but FMA-small has no such top genres, "
            "so those rows stay empty."
        ),
    },
    {
        "kicker": "Corpus",
        "title": "Official FMA-small, not a synthetic stand-in",
        "layout": "metrics",
        "metrics": [
            ("Clips", "511", "30 s  |  22 050 Hz"),
            ("Artists", "386", "no artist in two splits"),
            ("Test", "90", "held-out tracks"),
            ("Tags / clip", "6.24", "metadata-mapped"),
        ],
        "footnote": "MusicCaps audio and DEAM ratings are not on disk. Jazz / metal top-genres are unused.",
        "notes": (
            "We fetched FMA metadata and FMA-small from the official SWITCH mirror and SHA1-checked "
            "both zips. Seed forty-two, sixty-four clips from each of eight top genres. One mp3 "
            "failed to decode, leaving five hundred eleven tracks from three hundred eighty-six "
            "artists. Splits are artist-disjoint: three hundred twenty-seven, ninety-four, ninety. "
            "After mapping, folk and electronic have one hundred twenty-eight clips each; hip-hop "
            "has sixty-three. Experimental maps to electronic, instrumental to classical, "
            "international to folk. Captions are templates from title, artist, album, and leaf "
            "genres. MusicCaps CSV is on disk without YouTube audio. DEAM is gated, so valence "
            "and arousal are genre priors. Tags are keyword-mapped from those captions. That last "
            "fact will explain BERT's F1."
        ),
    },
    {
        "kicker": "Task 2",
        "title": "Two graphs, both deterministic functions of the spectrogram",
        "layout": "split",
        "split": [
            ("Segment graph", ["N = 30, 1 s windows", "Node dim 312", "Temporal + cosine > 0.62", "Self-loops kept"]),
            ("Chord graph", ["24 maj / min templates", "Collapse consecutive ids", "Edges = transition counts", "Mean N ~ 8.3"]),
        ],
        "footnote": "Dense batched GraphSAGE / GAT. No PyTorch Geometric.",
        "notes": (
            "The segment graph has thirty windows of one second. Each node concatenates mean and "
            "standard deviation of the audio streams: three hundred twelve dimensions. Adjacent "
            "windows get a temporal edge. If chroma concatenated with MFCC cosine exceeds tau 0.62, "
            "we add a similarity edge. That is a self-similarity matrix with a backbone. The chord "
            "graph matches frames to twenty-four major and minor triad templates, collapses runs, "
            "and weights edges by transition counts. Average size here is eight point three nodes. "
            "Neither graph is a musicological parse. Both give Task 2 a well-defined G. We save "
            "twenty-four example tensors for the lab."
        ),
    },
    {
        "kicker": "Tasks 1-2",
        "title": "TinyBERT from scratch, GraphSAGE vs a mel CNN",
        "layout": "bullets",
        "bullets": [
            "TinyBERT: 2 layers, d = 64, 4 heads, dropout 0.15, L = 48",
            "GraphSAGE: 2 layers, mean pool; GAT optional (2 heads, leaky 0.2)",
            "CNN: 16-32-64 channels on 64 x 32 log-mel",
            "AdamW 1.5e-3, batch 16, CPU, seed 42",
        ],
        "notes": (
            "TinyBERT is a two-layer pre-norm Transformer, hidden sixty-four, four heads. We train "
            "from scratch because five hundred eleven clips will not pretrain bert-base-uncased. "
            "The interface still matches HuggingFace CLS, so a real checkpoint can swap in later. "
            "GraphSAGE is two dense layers with mean pooling. GAT is implemented with two heads "
            "and an adjacency mask. The control is a three-block CNN on sixty-four by thirty-two "
            "log-mel. Same tag and genre heads. If the graph is doing extra work, GraphSAGE should "
            "beat the CNN. On eight-way genre, the CNN is slightly ahead: 0.489 versus 0.422."
        ),
    },
    {
        "kicker": "Tasks 3-4",
        "title": "Fuse g with the caption, then align two towers",
        "layout": "split",
        "split": [
            ("Fusion", ["Q = g W_Q, K,V from tokens", "Concat ablation: [g ; CLS]", "L = BCE + 0.35 VA L2"]),
            ("InfoNCE", ["Two 64-d spheres", "Symmetric batch loss", "R@K both directions"]),
        ],
        "notes": (
            "Fusion treats the graph vector as a query against caption keys and values, then an "
            "MLP on the concatenation. Early concat just stacks g with CLS. Three heads: tags, "
            "valence, arousal. Emotion loss weights alpha and beta are 0.35. BERT is warm-started "
            "from Task 1. Concat is the workhorse; cross-attention is the harder variant in the "
            "brief. Task 4 projects both towers onto a sixty-four-d sphere and trains symmetric "
            "InfoNCE at temperature 0.07. We score recall at one, five, and ten on the ninety "
            "pairs. Chance R at 5 is five over ninety, about 0.056."
        ),
    },
    {
        "kicker": "Held-out  |  N = 90",
        "title": "Text towers dominate tags; audio-only F1 is the honest number",
        "layout": "metrics",
        "metrics": [
            ("Concat F1", "0.783", "best tagger"),
            ("BERT F1", "0.772", "caption only"),
            ("GNN / CNN F1", "0.26", "audio only"),
            ("C to A R@5", "0.078", "chance 0.056"),
        ],
        "footnote": "AUC-PR concat 0.868  |  genre CNN 0.489 vs GNN 0.422  |  VA MAE 0.037 (priors, not DEAM).",
        "notes": (
            "Random tags: zero macro-F1. Majority: 0.118. CNN: 0.255. GNN: 0.264. BERT-only: 0.772. "
            "Concat fusion: 0.783, the best tagger, with AUC-PR 0.868. Cross-attention: 0.779, with "
            "a slightly better emotion MAE, 0.037 versus 0.039. Genre: CNN 0.489, GNN 0.422. Caption "
            "to audio R at 5 is 0.078, barely above chance 0.056. R at 1 is zero. Rank-1 is often a "
            "same-template neighbor, not the paired track. Look at the gap between the text tower "
            "and the audio tower. That gap is the result."
        ),
    },
    {
        "kicker": "Reading the table",
        "title": "BERT is decoding a template we wrote",
        "layout": "bullets",
        "bullets": [
            "Template leakage: tags come from caption metadata, so BERT / fusion F1 is inflated",
            "Audio-only ~0.26 F1 is the real FMA tagging difficulty here",
            "VA MAE 0.037 recovers genre priors, not listener ratings",
            "FMA captions are stereotyped, so InfoNCE cannot identify the pair",
        ],
        "notes": (
            "Why is BERT so strong? Because the twenty-four tags were extracted from the same "
            "metadata that forms the caption. Task 1 is closer to decoding a template than to "
            "open-vocabulary MusicCaps tagging. Fusion copies those tags almost exactly. That is "
            "expected, not impressive. Audio-only models at about 0.26 F1 are the honest FMA "
            "number at this capacity. Emotion MAE looks tiny because the targets are genre priors "
            "that both the caption and the graph can recover. It is not a DEAM human-rating result. "
            "Retrieval stays near chance because captions share the same bag of tokens. Two "
            "substitutions would change the table without touching the code: independent MusicCaps "
            "captions, and DEAM listener ratings. We report the FMA run as-is instead of simulating them."
        ),
    },
    {
        "kicker": "Error analysis",
        "title": "Genre collapse: electronic and hip-hop fall into folk",
        "layout": "split",
        "split": [
            ("000620 Arabesque", ["Caption: folk, Ed Askew", "Tags: copied", "Genre: folk to folk"]),
            ("007373 heartbreaker", ["Caption: electronic", "Tags: drops dark", "Genre: elec. to folk"]),
            ("013749 110% edit", ["Caption: hip-hop, Laws", "Tags: copied", "Genre: hip-hop to folk"]),
        ],
        "notes": (
            "Three held-out clips from the lab. Arabesque, folk, Ed Askew: fusion tags match, "
            "genre folk to folk. Heartbreaker, experimental electronic: the GNN predicts folk. A "
            "hip-hop edit by Laws: the GNN again predicts folk. Folk is the largest mapped class, "
            "one hundred twenty-eight clips, and short acoustic chroma is a convenient attractor. "
            "Thirty-second graphs at width sixty-four do not solve FMA-small genre. Retrieval also "
            "dies when two captions differ only in a title token that TinyBERT treats as unknown "
            "after the forty-eight-token cut."
        ),
    },
    {
        "kicker": "Takeaways",
        "title": "Four tasks on real FMA audio, with the caveats attached",
        "layout": "close",
        "bullets": [
            "Concat fusion 0.783 F1 — on labels built from captions",
            "GNN 0.264 / CNN 0.255 — honest audio tagging",
            "R@5 0.078 vs chance 0.056 — report the miss",
            "Lab: graphs, metrics, live fusion, IEEE paper",
        ],
        "footnote": "Questions  |  report/final_report.pdf  |  Cadence Lab",
        "notes": (
            "We implemented the full four-task pipeline on official FMA-small audio, with majority, "
            "CNN, BERT, and GNN baselines, concat versus cross-attention, retrieval, a seven-page "
            "IEEE report, twenty-four saved graphs, and the lab you can open after this talk. Concat "
            "fusion is the best tagger on metadata-derived labels. Audio-only models show the real "
            "difficulty. Retrieval needs unique captions. The remaining upgrade is MusicCaps audio "
            "and DEAM ratings, not a bigger Transformer on the same leaked tags. Live inference in "
            "the lab synthesizes an eight-second clip so the forward pass is visible without "
            "redistributing FMA mp3s. I am happy to take questions."
        ),
    },
]


def _set_run(run, *, size: int, color: RGBColor, bold: bool = False, font: str = "Calibri"):
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    run.font.name = font
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = rPr.find(qn(tag))
        if node is None:
            node = etree.SubElement(rPr, qn(tag))
        node.set("typeface", font)


def _fill(shape, color: RGBColor, line: RGBColor | None = None):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1)


def _textbox(slide, l, t, w, h, text, *, size, color, bold=False, font="Calibri", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run(run, size=size, color=color, bold=bold, font=font)
    return box


def _notes(slide, text: str):
    notes = slide.notes_slide
    notes.notes_text_frame.text = text


def _footer(slide, idx: int, n: int, footnote: str | None):
    _textbox(slide, Inches(0.55), Inches(7.18), Inches(10.2), Inches(0.28), footnote or "Cadence Lab", size=11, color=MUTE)
    _textbox(slide, Inches(11.2), Inches(7.18), Inches(1.6), Inches(0.28), f"{idx} / {n}", size=11, color=MUTE, align=PP_ALIGN.RIGHT)


def _header(slide, kicker: str, title: str, title_size: int = 32):
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.12), Inches(7.5))
    _fill(accent, COPPER)
    _textbox(slide, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.32), kicker.upper(), size=12, color=COPPER, bold=True)
    _textbox(slide, Inches(0.55), Inches(0.58), Inches(12.2), Inches(1.4), title, size=title_size, color=CREAM, bold=True, font="Georgia")


def _bg(slide):
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    _fill(bg, INK)
    spTree = slide.shapes._spTree
    sp = bg._element
    spTree.remove(sp)
    spTree.insert(2, sp)


def _card(slide, l, t, w, h, heading: str, items: list[str]):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, l, t, w, h)
    _fill(shape, PANEL, LINE)
    try:
        shape.adjustments[0] = 0.08
    except Exception:
        pass
    _textbox(slide, l + Inches(0.22), t + Inches(0.18), w - Inches(0.4), Inches(0.5), heading, size=16, color=COPPER, bold=True, font="Georgia")
    box = slide.shapes.add_textbox(l + Inches(0.22), t + Inches(0.72), w - Inches(0.4), h - Inches(0.9))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = "•  " + item
        _set_run(run, size=15, color=CREAM)


def _bullets(slide, items: list[str], *, y=Inches(2.2)):
    box = slide.shapes.add_textbox(Inches(0.7), y, Inches(12.0), Inches(4.7))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(16)
        run = p.add_run()
        run.text = "▸  " + item
        _set_run(run, size=22, color=CREAM)


def _metrics(slide, metrics: list[tuple[str, str, str]]):
    n = len(metrics)
    gap = Inches(0.22)
    left = Inches(0.55)
    usable = Inches(12.23)
    w = (usable - gap * (n - 1)) / n
    y = Inches(2.2)
    h = Inches(2.55)
    for i, (label, value, hint) in enumerate(metrics):
        x = left + i * (w + gap)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
        _fill(card, PANEL, LINE)
        try:
            card.adjustments[0] = 0.08
        except Exception:
            pass
        _textbox(slide, x + Inches(0.18), y + Inches(0.22), w - Inches(0.36), Inches(0.35), label.upper(), size=12, color=MUTE, bold=True)
        _textbox(slide, x + Inches(0.18), y + Inches(0.7), w - Inches(0.36), Inches(1.05), value, size=36, color=COPPER, bold=True, font="Georgia")
        _textbox(slide, x + Inches(0.18), y + Inches(1.9), w - Inches(0.36), Inches(0.4), hint, size=13, color=MUTE)


def _split_cards(slide, columns: list[tuple[str, list[str]]]):
    n = len(columns)
    gap = Inches(0.22)
    left = Inches(0.55)
    usable = Inches(12.23)
    w = (usable - gap * (n - 1)) / n
    y = Inches(2.15)
    h = Inches(4.55)
    for i, (heading, items) in enumerate(columns):
        x = left + i * (w + gap)
        _card(slide, x, y, w, h, heading, items)


def build() -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    n = len(SLIDES)

    for i, spec in enumerate(SLIDES, 1):
        slide = prs.slides.add_slide(blank)
        _bg(slide)
        title_size = 34 if spec["layout"] == "title" else 26
        _header(slide, spec["kicker"], spec["title"], title_size=title_size)

        if spec["layout"] in {"title", "bullets", "close"}:
            _bullets(slide, spec["bullets"])
        elif spec["layout"] == "metrics":
            _metrics(slide, spec["metrics"])
        elif spec["layout"] == "split":
            _split_cards(slide, spec["split"])

        _footer(slide, i, n, spec.get("footnote"))
        _notes(slide, spec["notes"])

    out = ROOT / "report" / "cadence_lab_slides.pptx"
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    public = ROOT / "web" / "public" / "cadence-lab-slides.pptx"
    public.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(out, public)
    return out


if __name__ == "__main__":
    path = build()
    print(f"wrote {path.relative_to(ROOT)}")
    print("lab copy → web/public/cadence-lab-slides.pptx")
